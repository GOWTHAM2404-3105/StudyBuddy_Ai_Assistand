"""Automated PowerPoint (.pptx) Generator for StudyBuddy.

Generates a 10-slide executive college project presentation with
custom widescreen (16:9) formatting, professional color palette,
metric tables, cards, and embedded presenter notes.
"""

import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.enum.shapes import MSO_SHAPE
except ImportError:
    print("[ERROR] python-pptx is not installed.")
    print("Please install it using: pip install python-pptx")
    sys.exit(1)

# Color Palette (Professional Blue-and-White Corporate Theme)
COLOR_PRIMARY = RGBColor(13, 110, 253)      # #0d6efd
COLOR_DARK = RGBColor(15, 23, 42)          # #0f172a
COLOR_MUTED = RGBColor(100, 116, 139)      # #64748b
COLOR_CARD_BG = RGBColor(248, 250, 252)    # #f8fafc
COLOR_WHITE = RGBColor(255, 255, 255)
COLOR_SUCCESS = RGBColor(16, 185, 129)     # #10b981
COLOR_ACCENT = RGBColor(2, 132, 199)       # #0284c7
COLOR_BORDER = RGBColor(226, 232, 240)     # #e2e8f0

prs = Presentation()
# Set widescreen 16:9 aspect ratio
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
blank_layout = prs.slide_layouts[6]


def add_slide_header(slide, title_text, category_text):
    """Add standardized professional header banner across slides."""
    # Category Tag
    cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8), Inches(0.35))
    tf_c = cat_box.text_frame
    tf_c.word_wrap = True
    p_c = tf_c.paragraphs[0]
    p_c.text = category_text.upper()
    p_c.font.size = Pt(11)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_PRIMARY

    # Title Text
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.8))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = title_text
    p_t.font.size = Pt(24)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_DARK


def set_speaker_notes(slide, notes_text):
    """Set presenter notes for the slide."""
    notes_slide = slide.notes_slide
    text_frame = notes_slide.notes_text_frame
    text_frame.text = notes_text


def create_card(slide, left, top, width, height, title, body_points, border_color=COLOR_BORDER):
    """Helper to create a structured card block."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLOR_CARD_BG
    shape.line.color.rgb = border_color
    shape.line.width = Pt(1.5)

    tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.2), width - Inches(0.4), height - Inches(0.4))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = title
    p0.font.bold = True
    p0.font.size = Pt(16)
    p0.font.color.rgb = COLOR_DARK
    p0.space_after = Pt(10)

    for pt in body_points:
        p = tf.add_paragraph()
        p.text = f"• {pt}"
        p.font.size = Pt(13)
        p.font.color.rgb = COLOR_MUTED
        p.space_after = Pt(6)


# ==============================================================================
# SLIDE 1: Title Slide
# ==============================================================================
s1 = prs.slides.add_slide(blank_layout)

# Hero Title Container
title_box = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(3.8))
tf1 = title_box.text_frame
tf1.word_wrap = True

p1 = tf1.paragraphs[0]
p1.text = "StudyBuddy"
p1.font.size = Pt(52)
p1.font.bold = True
p1.font.color.rgb = COLOR_PRIMARY

p2 = tf1.add_paragraph()
p2.text = "Student AI Assistant (AI + ML)"
p2.font.size = Pt(28)
p2.font.bold = True
p2.font.color.rgb = COLOR_DARK
p2.space_after = Pt(12)

p3 = tf1.add_paragraph()
p3.text = "An Intelligent College Study Companion Combining Local Generative AI (RAG)\n& Supervised Machine Learning Performance Forecasting"
p3.font.size = Pt(16)
p3.font.color.rgb = COLOR_MUTED
p3.space_after = Pt(24)

p4 = tf1.add_paragraph()
p4.text = "College Project Presentation  |  B.Tech / B.E. / MCA Technical Review"
p4.font.size = Pt(13)
p4.font.bold = True
p4.font.color.rgb = COLOR_ACCENT

set_speaker_notes(
    s1,
    "Good morning respected faculty and evaluators. Today I am presenting StudyBuddy—a functional, intermediate-level "
    "intelligent web application that combines local Generative AI, Retrieval-Augmented Generation, and Scikit-Learn "
    "machine learning pipelines to create an active, private study assistant for college students."
)

# ==============================================================================
# SLIDE 2: Problem Statement & Motivation
# ==============================================================================
s2 = prs.slides.add_slide(blank_layout)
add_slide_header(s2, "The Academic Challenge & Problem Statement", "Motivation")

create_card(
    s2, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "Information Overload",
    [
        "Students manage hundreds of pages of unindexed lecture slides and textbook PDFs.",
        "Manual keyword search fails on complex academic concepts and equations.",
        "Consuming large syllabi without structured summaries leads to study fatigue."
    ]
)

create_card(
    s2, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "Cloud AI Hallucinations & Privacy",
    [
        "Commercial cloud chatbots hallucinate facts without grounding or citations.",
        "Uploading confidential college assignments leaks student data to external servers.",
        "Cloud LLM API subscriptions are expensive and lack offline reliability."
    ]
)

create_card(
    s2, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "Late Academic Setbacks",
    [
        "Students discover conceptual gaps only after midterm or semester exam failures.",
        "Lack of early performance forecasting based on study time and attendance.",
        "Absence of formative, automated diagnostic testing before finals."
    ]
)

set_speaker_notes(
    s2,
    "We identified three core pain points in university education: static lecture slides causing information overload, "
    "cloud AI tools hallucinating answers without citations while leaking private student data, and a lack of early-warning "
    "systems to help students detect learning deficiencies before examination day."
)

# ==============================================================================
# SLIDE 3: Objectives & Vision
# ==============================================================================
s3 = prs.slides.add_slide(blank_layout)
add_slide_header(s3, "Project Objectives & Core Capabilities", "Objectives")

create_card(
    s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.3),
    "1. Privacy-First Local AI Tutor",
    [
        "Powered by locally hosted Ollama (Llama 3.2 / Qwen 2.5).",
        "Provides step-by-step pedagogical explanations without cloud fees."
    ]
)

create_card(
    s3, Inches(6.8), Inches(1.8), Inches(5.6), Inches(2.3),
    "2. Semantic RAG with Exact Citations",
    [
        "PyMuPDF per-page extraction and all-MiniLM-L6-v2 vector embeddings.",
        "Cites exact page numbers with strict anti-hallucination guardrails."
    ]
)

create_card(
    s3, Inches(0.8), Inches(4.4), Inches(5.6), Inches(2.3),
    "3. Active Recall AI Quiz Hub",
    [
        "Synthesizes 4-choice MCQs with server-separated answer keys.",
        "Instant scoring, mistaken option explanations, and retry capability."
    ]
)

create_card(
    s3, Inches(6.8), Inches(4.4), Inches(5.6), Inches(2.3),
    "4. ML Grade Prediction & Analytics",
    [
        "Random Forest Regressor trained on the UCI Student benchmark.",
        "Early forecast of final marks (G3) with Chart.js mastery telemetry."
    ]
)

set_speaker_notes(
    s3,
    "StudyBuddy's architecture rests on four pillars: a local AI tutor for conceptual explanations, a semantic RAG pipeline "
    "that grounds answers directly in uploaded PDFs, an automated quiz generator for active recall, and an ML model that "
    "provides early grade forecasting based on empirical educational data."
)

# ==============================================================================
# SLIDE 4: Architecture & Technology Stack
# ==============================================================================
s4 = prs.slides.add_slide(blank_layout)
add_slide_header(s4, "System Architecture & Technology Stack", "Architecture")

create_card(
    s4, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "Presentation Layer",
    [
        "HTML5, CSS3, JavaScript (ES6+)",
        "Responsive Bootstrap 5.3 UI framework",
        "Chart.js 4.4 for dynamic data charts",
        "Custom professional blue/white design",
        "Mobile-friendly collapsible navigation"
    ]
)

create_card(
    s4, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "Backend & Services",
    [
        "Python 3.13 & Flask 3.1 REST API",
        "Modular Blueprint architecture",
        "SQLAlchemy 2.0 ORM with SQLite DB",
        "PBKDF2:SHA256 password security",
        "Pytest automated testing framework"
    ]
)

create_card(
    s4, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "AI & Machine Learning",
    [
        "Ollama HTTP API (Llama 3.2 / Qwen)",
        "PyMuPDF (fitz) PDF parser",
        "Sentence Transformers (all-MiniLM-L6-v2)",
        "Cosine Similarity semantic retrieval",
        "Scikit-Learn Random Forest Pipeline"
    ]
)

set_speaker_notes(
    s4,
    "Here is our system architecture. We followed modern full-stack engineering standards: a clean RESTful Flask backend, "
    "SQLAlchemy ORM for data integrity, local AI inference through Ollama, and responsive Chart.js visual telemetry on the frontend."
)

# ==============================================================================
# SLIDE 5: Modules 1 & 2: Security & AI Chatbot
# ==============================================================================
s5 = prs.slides.add_slide(blank_layout)
add_slide_header(s5, "Module 1 & 2: Security & AI Study Chatbot", "Core Modules")

create_card(
    s5, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Module 1: Student Security & Isolation",
    [
        "Cryptographic password hashing via Werkzeug PBKDF2:SHA256.",
        "Custom @login_required decorator protecting all private routes.",
        "Strict multi-tenant database isolation: One student cannot view or query another student's documents, quizzes, or grade forecasts.",
        "Automatic session expiry and secure cookie management.",
        "Validated JSON API error handling with HTTP 401/409 codes."
    ]
)

create_card(
    s5, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Module 2: AI Study Chatbot Engine",
    [
        "Powered by local Ollama instance with zero external cloud dependencies.",
        "Pedagogical tutor persona: Explains concepts step-by-step with analogies and mathematical proofs.",
        "Multi-turn conversation memory maintained in SQLite.",
        "Frontend interactive bubble interface with real-time typing indicators.",
        "Resilient offline handling: Detects daemon state and guides the student."
    ]
)

set_speaker_notes(
    s5,
    "Module 1 guarantees student data privacy with cryptographic password hashing and database-enforced user isolation. "
    "Module 2 provides an academic tutor chatbot that maintains multi-turn conversation memory and breaks down complex questions step-by-step."
)

# ==============================================================================
# SLIDE 6: Module 3: PDF Notes Assistant & Semantic RAG
# ==============================================================================
s6 = prs.slides.add_slide(blank_layout)
add_slide_header(s6, "Module 3: PDF Notes Assistant (Semantic RAG)", "NLP & RAG")

create_card(
    s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Document Processing Pipeline",
    [
        "1. Ingestion: Upload lecture notes PDF (16MB max limit).",
        "2. Extraction: PyMuPDF parses text while tracking 1-indexed page numbers.",
        "3. Chunking: Sliding window of 500 characters with 100-character overlap preserves semantic context.",
        "4. Embeddings: Dense vectors generated via all-MiniLM-L6-v2.",
        "5. Retrieval: Normalized cosine similarity fetches top-4 most relevant context passages."
    ]
)

create_card(
    s6, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Grounded Q&A & Anti-Hallucination Guard",
    [
        "Strict Context Injection: The LLM is restricted to the retrieved passages.",
        "Exact Source Citations: Responses cite verifiable page numbers (e.g. [Source Page 2]).",
        "Hallucination Refusal: If notes do not contain the answer, the model states: 'The uploaded notes do not contain enough information to answer this question.'",
        "Executive Summaries: Automatically generates chapter summaries and bulleted key takeaways upon upload."
    ]
)

set_speaker_notes(
    s6,
    "Module 3 represents our semantic RAG implementation. Uploaded PDFs are parsed with PyMuPDF, chunked with sliding windows, "
    "and embedded via Sentence Transformers. When a student asks a question, the model cites exact page numbers and strictly "
    "refuses to hallucinate facts outside the provided document."
)

# ==============================================================================
# SLIDE 7: Module 4: AI Quiz Generator & Evaluation
# ==============================================================================
s7 = prs.slides.add_slide(blank_layout)
add_slide_header(s7, "Module 4: AI Quiz Generator & Evaluation Hub", "Active Recall")

create_card(
    s7, Inches(0.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "MCQ Question Synthesis",
    [
        "Generates 3 to 10 questions on any custom topic or uploaded PDF notes.",
        "Strict JSON schema enforcing exactly 4 distinct options (A, B, C, D).",
        "Every question generated includes an academic explanation.",
        "Fallback structured questions for reliable offline assessment."
    ]
)

create_card(
    s7, Inches(4.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "Anti-Cheating Architecture",
    [
        "Correct options are stored separately on the server database.",
        "Exam API serves only question text and options during test execution.",
        "Inspect-element attacks in the browser cannot reveal answer keys.",
        "Supports full retry capability for active recall reinforcement."
    ]
)

create_card(
    s7, Inches(8.8), Inches(1.8), Inches(3.6), Inches(4.8),
    "Auto-Scoring & Mistake Analysis",
    [
        "Instant automatic scoring upon submission with percentage calculation.",
        "Color-coded scorecard: Correct (+1) vs Incorrect choices.",
        "Detailed explanation rendered for every choice to clarify misconceptions.",
        "Logs attempt history to fuel the recommendation engine."
    ]
)

set_speaker_notes(
    s7,
    "Module 4 implements active recall testing. The quiz generator creates validated 4-choice questions. To uphold academic integrity, "
    "answers are stored on the server and hidden during the exam. After submission, answers are graded instantly with step-by-step explanations."
)

# ==============================================================================
# SLIDE 8: Module 5: ML Performance Prediction
# ==============================================================================
s8 = prs.slides.add_slide(blank_layout)
add_slide_header(s8, "Module 5: Machine Learning Grade Prediction", "Machine Learning")

# Benchmark Table
rows, cols = 3, 5
table_shape = s8.shapes.add_table(rows, cols, Inches(0.8), Inches(1.8), Inches(7.2), Inches(2.4))
table = table_shape.table

headers = ["Model Pipeline", "5-Fold CV RMSE", "Test MAE", "Test RMSE", "Test R²"]
for i, h in enumerate(headers):
    cell = table.cell(0, i)
    cell.text = h
    cell.fill.solid()
    cell.fill.fore_color.rgb = COLOR_DARK
    p = cell.text_frame.paragraphs[0]
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = COLOR_WHITE

data_rows = [
    ["Linear Regression (Baseline)", "4.3030", "3.5606", "4.4581", "0.0307"],
    ["Random Forest Regressor", "4.1317", "3.4194", "4.2867", "0.1038"]
]

for row_idx, r_data in enumerate(data_rows, start=1):
    for col_idx, val in enumerate(r_data):
        cell = table.cell(row_idx, col_idx)
        cell.text = val
        cell.fill.solid()
        cell.fill.fore_color.rgb = COLOR_CARD_BG if row_idx == 1 else RGBColor(240, 253, 244)
        p = cell.text_frame.paragraphs[0]
        p.font.size = Pt(11)
        if row_idx == 2:
            p.font.bold = True
            p.font.color.rgb = COLOR_SUCCESS

# Features Card
create_card(
    s8, Inches(8.3), Inches(1.8), Inches(4.2), Inches(4.8),
    "Feature Importance Weights",
    [
        "Lecture Absences (absences): 47.39%\nStrongest non-linear indicator of study consistency.",
        "Past Class Failures (failures): 39.98%\nSignificant historical drag on predicted final marks.",
        "Weekly Study Time (studytime): 12.63%\nPositive factor (>5 hrs/week yields 2-3 pt boost).",
        "Dataset: UCI Student Performance (student-mat.csv, N=395).",
        "Strict early prediction: G1 and G2 excluded."
    ]
)

# Textbox below table
tb_exp = s8.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(7.2), Inches(2.1))
tf_exp = tb_exp.text_frame
tf_exp.word_wrap = True
p_e = tf_exp.paragraphs[0]
p_e.text = "Experimental Setup & Ethical Compliance:"
p_e.font.bold = True
p_e.font.size = Pt(13)
p_e.font.color.rgb = COLOR_DARK

pts_e = [
    "Predicts final grade G3 on a 0–20 scale early in the academic term.",
    "Trained pipeline (StandardScaler + RandomForest) serialized using Joblib.",
    "Includes prominent statistical disclaimer: Estimates serve as early advisory signals, not deterministic exam guarantees."
]
for pt in pts_e:
    p = tf_exp.add_paragraph()
    p.text = f"• {pt}"
    p.font.size = Pt(11.5)
    p.font.color.rgb = COLOR_MUTED

set_speaker_notes(
    s8,
    "In Module 5, we trained an early grade predictor on the UCI Student Performance dataset. To avoid data leakage, "
    "we strictly excluded intermediate term marks G1 and G2. Random Forest outperformed Linear Regression with an RMSE "
    "of 4.28. The top predictive features were absences at 47% and past failures at 40%."
)

# ==============================================================================
# SLIDE 9: Modules 6 & 7: Recommendations & Analytics
# ==============================================================================
s9 = prs.slides.add_slide(blank_layout)
add_slide_header(s9, "Modules 6 & 7: Recommendations & Dashboard Analytics", "Analytics")

create_card(
    s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Module 6: Personalized Recommendations",
    [
        "Transparent Rule-Based Heuristic (clearly disclosed without false ML claims).",
        "Topic Deficiency Categorization:\n  - Score < 60%: Critical Review (Red Badge)\n  - Score 60-74%: Practice Recommended (Yellow)\n  - Score ≥ 75%: Mastered (Green Badge)",
        "Contextual Action Items: Prescribes retry quizzes, deep-dive AI tutor prompts, and direct links to matching uploaded notes.",
        "Extensible Architecture: Ready for collaborative filtering ML once student interaction logs expand."
    ]
)

create_card(
    s9, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Module 7: Student Analytics Dashboard",
    [
        "Telemetry Summary Cards: Total quizzes taken, average score, uploaded notes count, latest predicted GPA.",
        "Chart.js Topic Mastery Bar Chart: Compares student scores across attempted curriculum topics.",
        "Chart.js Progress Line Chart: Tracks student score trajectory over time to visualize learning velocity.",
        "Recent Quiz Results Feed with quick breakdown links.",
        "Quick action shortcuts for chatbot, PDF uploads, and grade forecasting."
    ]
)

set_speaker_notes(
    s9,
    "Module 6 maps diagnostic quiz scores to targeted study interventions, identifying weak topics and linking them "
    "to relevant lecture slides. Module 7 visualizes this telemetry using interactive Chart.js bar and line graphs, "
    "giving students an intuitive overview of their academic standing."
)

# ==============================================================================
# SLIDE 10: Technical Validation, Future Scope & Conclusion
# ==============================================================================
s10 = prs.slides.add_slide(blank_layout)
add_slide_header(s10, "Technical Validation, Future Scope & Conclusion", "Conclusion")

create_card(
    s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Technical Validation & Quality Assurance",
    [
        "100% Automated Test Pass Rate (14/14 tests in Pytest):",
        "  ✓ test_auth.py: Registration, PBKDF2 hashing, duplicate checks.",
        "  ✓ test_ml_predict.py: Bounds validation, Joblib pipeline integrity.",
        "  ✓ test_quiz.py: 4-choice schema, hidden answers, auto-scoring.",
        "  ✓ test_rag.py: PyMuPDF extraction, sliding chunking, cosine retrieval.",
        "Strict input validation and boundary protections across all REST APIs.",
        "Zero external cloud API costs or data leakage risks."
    ]
)

create_card(
    s10, Inches(6.8), Inches(1.8), Inches(5.6), Inches(4.8),
    "Future Scope & Conclusion",
    [
        "Future Enhancements:",
        "  • Multi-modal diagram understanding from textbook figures.",
        "  • Audio lecture transcription via local Whisper models.",
        "  • Trained collaborative filtering recommendation models.",
        "Conclusion:",
        "  StudyBuddy demonstrates that genuine Artificial Intelligence and Machine Learning can be integrated into an accessible, privacy-first college study platform.",
        "Thank you! Questions & Discussion."
    ]
)

set_speaker_notes(
    s10,
    "To conclude, StudyBuddy is completely built, validated with 14 automated unit tests, and ready for deployment. "
    "It demonstrates the power of combining local LLMs, semantic RAG, and classical machine learning into a unified, "
    "privacy-preserving college study assistant. Thank you, and I look forward to your questions."
)

# ==============================================================================
# Output Presentation File
# ==============================================================================
output_dir = Path(__file__).resolve().parent
output_file = output_dir / "StudyBuddy_Presentation.pptx"
prs.save(str(output_file))
print(f"[OK] PowerPoint presentation generated successfully at:\n     {output_file}")
